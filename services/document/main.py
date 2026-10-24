"""Offline document extraction. No arbitrary paths/URLs; bounded uploads and pages."""
import base64
import io
import os
import zipfile
from collections import Counter
from fastapi import FastAPI, UploadFile, HTTPException
import pymupdf as fitz
from docx import Document
from PIL import Image, ImageChops, ImageStat
import pytesseract

app = FastAPI(title="Resume document engine")
Image.MAX_IMAGE_PIXELS = 25_000_000
MAX_BYTES = 15 * 1024 * 1024
MAX_PAGES = 12

def png64(image):
    image.thumbnail((1000, 1415))
    out = io.BytesIO()
    image.convert("RGB").save(out, format="PNG")
    return base64.b64encode(out.getvalue()).decode()

def render_page(page):
    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
    return Image.open(io.BytesIO(pix.tobytes("png")))

def ocr(image):
    return pytesseract.image_to_string(image, lang=os.getenv("OCR_LANGUAGES", "eng+tha"), timeout=40)

def extract(data, name):
    lower = name.lower()
    if lower.endswith(".pdf"):
        doc = fitz.open(stream=data, filetype="pdf")
        if doc.is_encrypted:
            raise ValueError("Password-protected PDFs must be unlocked first")
        if len(doc) > MAX_PAGES:
            raise ValueError(f"Maximum {MAX_PAGES} pages")
        pages, texts, previews, warnings = [], [], [], []
        for number, page in enumerate(doc):
            if page.rect.width * page.rect.height > 2_000_000:
                raise ValueError("Page dimensions too large")
            text = page.get_text(sort=True)
            image = render_page(page)
            if len(text.strip()) < 30:
                text = ocr(image)
                warnings.append(f"Page {number+1}: OCR text must be reviewed")
            texts.append(text)
            if number < 2:
                previews.append(png64(image))
            spans = []
            for block in page.get_text("dict", flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)["blocks"]:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        spans.append({k: span[k] for k in ("text", "font", "size", "flags", "color", "bbox")})
            drawings = [{"rect": list(d["rect"]), "fill": d.get("fill"), "color": d.get("color"), "width": d.get("width")} for d in page.get_drawings()[:100]]
            pages.append({"page": number+1, "width": page.rect.width, "height": page.rect.height, "spans": spans[:800], "drawings": drawings, "images": [{"bbox": list(x["bbox"]), "width": x["width"], "height": x["height"]} for x in page.get_image_info()[:20]]})
        doc.close()
        return {"text": "\n\n".join(texts)[:60000], "previews": previews, "layout": {"pages": pages}, "warnings": warnings + ["Fonts may be substituted. Embedded photos, icons and decorative graphics are recorded but not automatically rebuilt."]}
    if lower.endswith(".docx"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 40*1024*1024 or len(z.infolist()) > 2000:
                raise ValueError("DOCX expanded size exceeds limit")
        doc = Document(io.BytesIO(data))
        text = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                text.append(" | ".join(c.text for c in row.cells))
        paragraphs = [{"text": p.text, "style": p.style.name, "alignment": str(p.alignment), "runs": [{"text": r.text, "font": r.font.name, "size": r.font.size.pt if r.font.size else None, "bold": r.bold, "color": str(r.font.color.rgb) if r.font.color.rgb else None} for r in p.runs]} for p in doc.paragraphs[:300]]
        page = doc.sections[0]
        return {"text": "\n".join(text)[:60000], "previews": [], "layout": {"width": page.page_width.pt, "height": page.page_height.pt, "margin": page.left_margin.mm, "paragraphs": paragraphs, "tables": len(doc.tables)}, "warnings": ["DOCX design uses OOXML text styles. Floating shapes/photos are not automatically rebuilt; no original-page preview is available."]}
    if lower.endswith((".png", ".jpg", ".jpeg")):
        image = Image.open(io.BytesIO(data)).convert("RGB")
        return {"text": ocr(image)[:60000], "previews": [png64(image)], "layout": {"imageWidth": image.width, "imageHeight": image.height}, "warnings": ["Image-based layout and OCR are approximate. Confirm every fact."]}
    raise ValueError("Unsupported file type")

async def upload_bytes(file):
    data = await file.read(MAX_BYTES+1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "Maximum upload is 15 MB")
    return data

@app.get("/health")
def health():
    return {"ok": True}

@app.post("/extract")
def extract_route(file: UploadFile):
    data = file.file.read(MAX_BYTES+1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "Maximum upload is 15 MB")
    try:
        return extract(data, file.filename or "")
    except Exception as e:
        raise HTTPException(422, str(e))

@app.post("/inspect")
def inspect(file: UploadFile):
    data = file.file.read(MAX_BYTES+1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413)
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        if len(doc)>MAX_PAGES: raise ValueError("Too many pages")
        return {"pages": len(doc), "text": "\n".join(p.get_text() for p in doc)}
    except Exception as e:
        raise HTTPException(422, str(e))

@app.post("/compare")
def compare(original: UploadFile, generated: UploadFile):
    try:
        a = original.file.read(MAX_BYTES+1)
        b = generated.file.read(MAX_BYTES+1)
        if max(len(a), len(b))>MAX_BYTES: raise ValueError("File too large")
        source = extract(a, original.filename or "")
        if not source["previews"]: raise ValueError("Visual comparison requires PDF or image original")
        target = extract(b, "generated.pdf")
        x = Image.open(io.BytesIO(base64.b64decode(source["previews"][0]))).convert("RGB")
        y = Image.open(io.BytesIO(base64.b64decode(target["previews"][0]))).convert("RGB").resize(x.size)
        difference = ImageChops.difference(x, y)
        mae = sum(ImageStat.Stat(difference).mean)/3
        return {"original": source["previews"][0], "generated": target["previews"][0], "difference": png64(difference), "pixelAgreement": round(100*(1-mae/255), 2), "metric": "First-page normalized pixel MAE; whitespace and changed content affect the value. NOT an ATS, accuracy or layout-quality score."}
    except Exception as e:
        raise HTTPException(422, str(e))
