import io
import unittest
import fitz
from docx import Document
from fastapi.testclient import TestClient
from main import app, extract
from PIL import Image

class DocumentTests(unittest.TestCase):
    def setUp(self):
        d = fitz.open()
        p = d.new_page()
        p.insert_text((40, 50), "EXAMPLE PERSON", fontsize=22)
        p.insert_text((40, 90), "Backend Developer. Built Node.js APIs for ten projects.")
        p.draw_rect(fitz.Rect(20, 110, 150, 780), fill=(0.1, 0.3, 0.4))
        self.pdf = d.tobytes()
        self.client = TestClient(app)
    def test_pdf_geometry_text_color_and_preview(self):
        r = extract(self.pdf, 'cv.pdf')
        self.assertIn('Backend Developer', r['text'])
        self.assertIn('Backend Developer', r['textPages'][0])
        self.assertTrue(r['previews'])
        self.assertTrue(r['layout']['pages'][0]['drawings'])
        self.assertEqual(r['layout']['pages'][0]['spans'][0]['size'], 22)
    def test_embedded_header_photo_is_recovered(self):
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((40,60), "Example Engineer with backend development experience and document preview testing")
        image = Image.new('RGB',(100,100),(28,85,120))
        b = io.BytesIO()
        image.save(b,format='PNG')
        p.insert_image(fitz.Rect(470,35,545,110),stream=b.getvalue())
        data = extract(doc.tobytes(),'profile.pdf')
        self.assertIsInstance(data['profilePhoto'],str)
        self.assertGreater(len(data['profilePhoto']),100)
    def test_docx_styles_and_tables(self):
        d = Document()
        d.add_paragraph('Actual developer', 'Title')
        table = d.add_table(rows=1, cols=2)
        table.cell(0,0).text='React'
        table.cell(0,1).text='Personal project'
        stream = io.BytesIO();d.save(stream)
        r=extract(stream.getvalue(), 'cv.docx')
        self.assertIn('Personal project', r['text'])
        self.assertEqual(r['layout']['tables'],1)
    def test_page_limit(self):
        d=fitz.open()
        for i in range(13): d.new_page()
        with self.assertRaisesRegex(ValueError,'Maximum'): extract(d.tobytes(),'x.pdf')
    def test_invalid_upload(self):
        r=self.client.post('/extract',files={'file':('cv.pdf',b'invalid')})
        self.assertEqual(r.status_code,422)
    def test_identical_visual_comparison(self):
        r=self.client.post('/compare',files={'original':('a.pdf',self.pdf),'generated':('b.pdf',self.pdf)})
        self.assertEqual(r.status_code,200)
        self.assertEqual(r.json()['pixelAgreement'],100)
    def test_inspect_selectable_text(self):
        r=self.client.post('/inspect',files={'file':('cv.pdf',self.pdf)})
        self.assertEqual(r.json()['pages'],1)
        self.assertIn('EXAMPLE PERSON',r.json()['text'])
if __name__=='__main__': unittest.main()
